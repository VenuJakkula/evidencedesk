"""Bounded XLSX parsing. No formula evaluation, URL access or archive extraction."""
import io
import os
import re
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath

from defusedxml import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from pydantic import BaseModel, ConfigDict, Field, StrictBool


class Problem(Exception):
    def __init__(self, code, message, status=422):
        self.code, self.message, self.status = code, message, status
        super().__init__(message)


@dataclass(frozen=True)
class Limits:
    upload: int = 10 * 1024 * 1024
    expanded: int = 100 * 1024 * 1024
    sheets: int = 20
    questions: int = 5000
    cells: int = 100000
    rows: int = 20000
    entries: int = 1000

    @classmethod
    def configured(cls):
        values = {name: int(os.getenv('EVIDENCEDESK_MAX_' + name.upper(), default))
                  for name, default in cls().__dict__.items()}
        if any(value < 1 for value in values.values()):
            raise RuntimeError('Workbook limits must be positive')
        return cls(**values)


class SheetMapping(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    sheet: str = Field(min_length=1, max_length=31)
    header_row: int = Field(ge=1, le=20000)
    question_column: str = Field(pattern=r'^[A-Z]{1,3}$')
    answer_column: str = Field(pattern=r'^[A-Z]{1,3}$')
    identifier_column: str | None = Field(default=None, pattern=r'^[A-Z]{1,3}$')
    include_hidden: StrictBool = False
    overwrite_existing: StrictBool = False


class Mapping(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    sheets: list[SheetMapping] = Field(min_length=1, max_length=20)


def safe_filename(filename):
    name = (filename or '').replace('\\', '/').split('/')[-1]
    name = re.sub(r'[^A-Za-z0-9_. -]', '_', name)[:150]
    return name if name.lower().endswith('.xlsx') else 'workbook.xlsx'


def workbook(data, filename, limits=None):
    limits = limits or Limits.configured()
    if len(data) > limits.upload:
        raise Problem('upload_too_large', 'Workbook exceeds the configured upload limit.', 413)
    if not filename or not filename.lower().endswith('.xlsx'):
        raise Problem('unsupported_format', 'Upload an unencrypted .xlsx workbook; .xls and .xlsm are unsupported.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.entries or sum(e.file_size for e in entries) > limits.expanded:
                raise Problem('archive_too_large', 'Expanded workbook exceeds archive limits.', 413)
            names = [e.filename for e in entries]
            if len(set(names)) != len(names):
                raise Problem('malformed_workbook', 'Duplicate archive entries are unsupported.')
            if not {'[Content_Types].xml', 'xl/workbook.xml'} <= set(names):
                raise Problem('malformed_workbook', 'Missing XLSX workbook structure.')
            cell_count = sheet_count = merged_cells = 0
            for entry in entries:
                name = entry.filename
                path = PurePosixPath(name)
                if path.is_absolute() or '..' in path.parts or '\\' in name or ':' in name:
                    raise Problem('unsafe_archive', 'Workbook archive contains an unsafe entry name.')
                if entry.flag_bits & 1:
                    raise Problem('unsupported_format', 'Encrypted workbooks are unsupported.')
                # This explicit part allowlist excludes VBA, objects, charts, drawings,
                # external links, comments, pivots, tables and custom XML extensions.
                if not re.fullmatch(r'(\[Content_Types\]\.xml|_rels/\.rels|docProps/(core|app)\.xml|xl/(workbook\.xml|styles\.xml|sharedStrings\.xml|theme/theme[0-9]+\.xml|_rels/workbook\.xml\.rels|worksheets/sheet[0-9]+\.xml|worksheets/_rels/sheet[0-9]+\.xml\.rels))', name):
                    raise Problem('unsupported_feature', 'Workbook contains unsupported objects or package parts.')
                root = ET.fromstring(archive.read(entry))
                if name == '[Content_Types].xml':
                    types = [node.get('ContentType', '') for node in root]
                    if 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml' not in types:
                        raise Problem('unsupported_format', 'Only ordinary .xlsx workbooks are supported; macros and templates are unsupported.')
                for node in root.iter():
                    tag = node.tag.rsplit('}', 1)[-1]
                    if tag in {'extLst', 'AlternateContent', 'oleObjects', 'drawing', 'legacyDrawing'}:
                        raise Problem('unsupported_feature', 'Workbook contains unsupported extensions or objects.')
                    if tag == 'Relationship' and (node.get('TargetMode') == 'External' or 'externalLink' in node.get('Type', '')):
                        raise Problem('external_links', 'External relationships are unsupported; remove links before import.')
                    if tag == 'sheet':
                        sheet_count += 1
                    if name.startswith('xl/worksheets/') and tag in {'c', 'row', 'dimension', 'mergeCell'}:
                        reference = node.get('r') or node.get('ref', '')
                        for number in re.findall(r'\d+', reference):
                            if int(number) > limits.rows:
                                raise Problem('worksheet_too_large', 'Worksheet row extent exceeds the configured limit.', 413)
                        if tag == 'c':
                            cell_count += 1
                        if tag == 'mergeCell':
                            left, top, right, bottom = range_boundaries(reference)
                            merged_cells += (right - left + 1) * (bottom - top + 1)
                            if merged_cells > limits.cells:
                                raise Problem('worksheet_too_large', 'Merged-cell expansion exceeds populated-cell limits.', 413)
                if sheet_count > limits.sheets or cell_count > limits.cells:
                    raise Problem('worksheet_too_large', 'Workbook exceeds sheet or populated-cell limits.', 413)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            result = load_workbook(io.BytesIO(data), data_only=False, keep_links=False, rich_text=True)
            if caught:
                result.close()
                raise Problem('unsupported_feature', 'Workbook includes features the XLSX library cannot preserve.')
        return result
    except Problem:
        raise
    except Exception as exc:
        raise Problem('malformed_workbook', 'Invalid, encrypted or malformed XLSX workbook.') from exc


def inspect(data, filename, mapping=None, limits=None):
    limits = limits or Limits.configured()
    wb = workbook(data, filename, limits)
    try:
        sheets = [{'name': ws.title, 'state': ws.sheet_state, 'rows': ws.max_row,
                   'candidate_columns': [{'column': get_column_letter(col),
                                          'sample': str(ws.cell(1, col).value or '')[:200]}
                                         for col in range(1, min(ws.max_column, 100) + 1)]}
                  for ws in wb]
        if mapping is None:
            return {'sheets': sheets, 'mapping': None, 'rows': [], 'total': 0,
                    'warnings': ['Select sheets, header rows and columns explicitly before importing.']}
        rows, messages, destinations, questions = [], [], set(), set()
        for spec in mapping.sheets:
            if spec.question_column == spec.answer_column or spec.identifier_column in {spec.question_column, spec.answer_column}:
                raise Problem('invalid_mapping', 'Choose distinct question, answer and identifier columns.')
            for other in mapping.sheets:
                if other is not spec and spec.sheet == other.sheet and spec.answer_column in {other.question_column, other.identifier_column}:
                    raise Problem('overlapping_mapping', 'Answer mappings overlap question or identifier columns.')
        for spec in mapping.sheets:
            if spec.sheet not in wb.sheetnames:
                raise Problem('invalid_mapping', 'Mapped sheet does not exist.')
            ws = wb[spec.sheet]
            cols = [column_index_from_string(c) for c in (spec.question_column, spec.answer_column, spec.identifier_column) if c]
            if max(cols) > 16384 or len(cols) != len(set(cols)) or spec.header_row > ws.max_row:
                raise Problem('invalid_mapping', 'Choose distinct valid question, answer and identifier columns and an existing header row.')
            if ws.sheet_state != 'visible' and not spec.include_hidden:
                raise Problem('hidden_sheet', 'Explicit include_hidden intent is required for hidden sheets.')
            for row in range(spec.header_row + 1, ws.max_row + 1):
                if ws.row_dimensions[row].hidden and not spec.include_hidden:
                    messages.append({'code': 'hidden_row_skipped', 'sheet': ws.title, 'row': row})
                    continue
                question = ws[f'{spec.question_column}{row}']
                answer = ws[f'{spec.answer_column}{row}']
                if question.value is None or not str(question.value).strip():
                    continue
                if question.data_type == 'f':
                    raise Problem('formula_question', 'Formula question cells are unsupported; use literal question text.')
                if not isinstance(question.value, str) or len(question.value) > 5000:
                    raise Problem('invalid_question', 'Questions must be literal text of at most 5000 characters.')
                if isinstance(answer, MergedCell):
                    raise Problem('merged_destination', 'Answer cells must be merged-range anchors or unmerged cells.')
                if answer.data_type == 'f':
                    raise Problem('formula_destination', 'Formula answer cells cannot be overwritten.')
                dest = (ws.title, answer.coordinate)
                source = (ws.title, question.coordinate)
                if dest in destinations:
                    raise Problem('duplicate_destination', 'Multiple mappings select the same answer cell.')
                destinations.add(dest)
                questions.add(source)
                existing = answer.value is not None and answer.value != ''
                writable = not existing or spec.overwrite_existing
                if existing:
                    messages.append({'code': 'existing_answer_overwrite' if writable else 'existing_answer_preserved',
                                     'sheet': ws.title, 'cell': answer.coordinate})
                identifier = ws[f'{spec.identifier_column}{row}'].value if spec.identifier_column else None
                rows.append({'sheet': ws.title, 'row': row, 'question_cell': question.coordinate,
                             'answer_cell': answer.coordinate, 'source_identifier': str(identifier) if identifier is not None else None,
                             'question': question.value, 'writable': writable,
                             'overwrite_existing': spec.overwrite_existing})
                if len(rows) > limits.questions:
                    raise Problem('too_many_questions', 'Import exceeds the configured question limit.', 413)
        if questions & destinations:
            raise Problem('overlapping_mapping', 'Question and answer mappings overlap.')
        # Include merged anchor footprints: writing an anchor must not consume another source/destination.
        for ws in wb:
            for merged in ws.merged_cells.ranges:
                anchor = f'{get_column_letter(merged.min_col)}{merged.min_row}'
                if (ws.title, anchor) in destinations:
                    for sheet, coordinate in questions | destinations:
                        if sheet == ws.title and coordinate != anchor and coordinate in merged:
                            raise Problem('overlapping_mapping', 'Merged destination overlaps another mapped cell.')
        if not rows:
            raise Problem('empty_questionnaire', 'Mapping contains no nonempty questions.')
        return {'sheets': sheets, 'mapping': mapping.model_dump(), 'rows': rows,
                'total': len(rows), 'warnings': messages}
    finally:
        wb.close()
