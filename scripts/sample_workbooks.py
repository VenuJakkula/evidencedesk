"""Generate two synthetic supported templates: python -m scripts.sample_workbooks."""
import io
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


def template(kind='security'):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Security' if kind == 'security' else 'RFP'
    header = 1 if kind == 'security' else 3
    if header == 3:
        ws['A1'] = 'Synthetic procurement questionnaire'
        ws.merge_cells('A1:D1')
    for col, value in enumerate(['ID', 'Question', 'Answer', 'Notes'], 1):
        ws.cell(header, col, value)
    questions = ['Is customer data encrypted at rest?', 'Is customer data encrypted at rest?',
                 None, 'What transport encryption is used?', 'Is biometric authentication available?',
                 'Is customer data encrypted at rest?']
    for index, question in enumerate(questions, header + 1):
        ws.cell(index, 1, f'Q{index - header}')
        ws.cell(index, 2, question)
    ws.cell(header + 4, 3, 'Existing response')
    ws.cell(header + 6, 3, 'Hidden existing response')
    ws.row_dimensions[header + 6].hidden = True
    ws['F2'] = '=1+2'
    ws['G2'] = 'Styled non-target cell'
    ws['G2'].font = Font(name='Calibri', bold=True, color='FF123456')
    ws['G2'].fill = PatternFill('solid', fgColor='FFABCDEF')
    ws['G2'].alignment = Alignment(wrap_text=True, horizontal='center')
    ws['G2'].border = Border(bottom=Side(style='thin', color='FF000000'))
    ws.column_dimensions['B'].width = 48
    ws.column_dimensions['C'].width = 60
    ws.row_dimensions[header + 1].height = 35
    ws.merge_cells(f'C{header + 2}:D{header + 2}')
    validation = DataValidation(type='list', formula1='"Yes,No,Unknown"')
    ws.add_data_validation(validation)
    validation.add('H2:H5')
    ws.freeze_panes = f'C{header + 1}'
    second = wb.create_sheet('Operations')
    second.append(['Question', 'Answer'])
    second.append(['Is customer data encrypted at rest?', None])
    second['D4'] = '=SUM(1,2)'
    hidden = wb.create_sheet('Internal')
    hidden.append(['Question', 'Answer'])
    hidden.append(['Is customer data encrypted at rest?', None])
    hidden.sheet_state = 'hidden'
    mapping = {'sheets': [
        {'sheet': ws.title, 'header_row': header, 'question_column': 'B', 'answer_column': 'C', 'identifier_column': 'A'},
        {'sheet': 'Operations', 'header_row': 1, 'question_column': 'A', 'answer_column': 'B'}]}
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue(), mapping


def main():
    directory = Path(__file__).resolve().parents[1] / 'samples'
    directory.mkdir(exist_ok=True)
    for kind in ('security', 'rfp'):
        data, mapping = template(kind)
        (directory / f'{kind}.xlsx').write_bytes(data)
        (directory / f'{kind}.mapping.json').write_text(json.dumps(mapping, indent=2) + '\n')
    print('Generated two synthetic workbooks and explicit mapping files in samples/.')


if __name__ == '__main__':
    main()
