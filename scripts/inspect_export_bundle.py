import argparse
import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook


def verify(bundle_path, expected):
    """Inspect an actual browser download using the backend's existing XLSX library."""
    source = Path(__file__).resolve().parents[1] / 'samples' / (expected['kind'] + '.xlsx')
    with zipfile.ZipFile(bundle_path) as archive:
        assert set(archive.namelist()) == {'completed.xlsx', 'unresolved-items.csv', 'manifest.json'}
        output = archive.read('completed.xlsx')
        manifest = json.loads(archive.read('manifest.json'))
        report = list(csv.DictReader(io.StringIO(archive.read('unresolved-items.csv').decode('utf-8'))))
    selected = expected['selected']
    actual_cells = {f"{row['sheet']}!{row['cell']}" for row in manifest['selected_revisions']}
    assert actual_cells == set(selected), (actual_cells, set(selected))
    assert len(report) == 5 - len(selected)
    assert len(manifest['unresolved_item_ids']) == len(report)
    assert {row['item_id'] for row in report} == set(manifest['unresolved_item_ids'])
    assert manifest['source_sha256'] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert manifest['output_sha256'] == hashlib.sha256(output).hexdigest()
    original = load_workbook(source, data_only=False)
    completed = load_workbook(io.BytesIO(output), data_only=False)
    try:
        assert original.sheetnames == completed.sheetnames
        for original_sheet in original:
            sheet = completed[original_sheet.title]
            assert sheet.sheet_state == original_sheet.sheet_state
            assert set(str(value) for value in sheet.merged_cells.ranges) == set(str(value) for value in original_sheet.merged_cells.ranges)
            for row in original_sheet:
                for cell in row:
                    address = f'{original_sheet.title}!{cell.coordinate}'
                    target = sheet[cell.coordinate]
                    if address in selected:
                        assert target.value == selected[address], (address, target.value)
                        assert target.data_type == 's'
                    else:
                        assert target.value == cell.value, (address, target.value, cell.value)
                        if cell.data_type == 'f':
                            assert target.data_type == 'f'
        for cell, reason in expected.get('unresolved', {}).items():
            assert any(f"{row['sheet']}!{row['answer_cell']}" == cell and row['reason'] == reason for row in report), (cell, reason)
    finally:
        original.close(); completed.close()
    return {'selected_count': len(selected), 'unresolved_count': len(report), 'export_id': manifest['export_id'], 'unchanged_cells_verified': True}


def main():
    parser = argparse.ArgumentParser(description='Verify a synthetic browser export against its sample workbook.')
    parser.add_argument('bundle')
    parser.add_argument('expected_json')
    args = parser.parse_args()
    print(json.dumps(verify(args.bundle, json.loads(args.expected_json))))


if __name__ == '__main__':
    main()
