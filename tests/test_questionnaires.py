import hashlib
import io
import json
import sqlite3
import struct
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.core import Store
from app.main import app, store as api_store, tenant
from app.questionnaires import Questionnaires
from app.workbooks import Limits, Mapping, Problem, inspect, workbook
from scripts.sample_workbooks import template


@pytest.fixture
def service(tmp_path):
    store = Store(tmp_path / 'workflow.db')
    store.add_document('a', 'Policy', 'Customer data is encrypted at rest using AES-256. Transport encryption uses TLS 1.2 or higher.', '1', '2099-01-01')
    return Questionnaires(store)


def imported(service, kind='security', key='key', raw_mapping=None):
    data, mapping = template(kind)
    result = service.create('a', data, f'{kind}.xlsx', Mapping.model_validate(raw_mapping or mapping), key)
    return result, data, mapping


def unpack(bundle):
    with zipfile.ZipFile(io.BytesIO(bundle)) as archive:
        return archive.read('completed.xlsx'), json.loads(archive.read('manifest.json')), archive.read('unresolved-items.csv').decode()


def mutate(data, change):
    wb = load_workbook(io.BytesIO(data))
    change(wb)
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def error(code, call):
    with pytest.raises(Problem) as exc:
        call()
    assert exc.value.code == code


@pytest.mark.parametrize('kind,header', [('security', 1), ('rfp', 3)])
def test_import_exact_mapping_preview_and_restart(service, kind, header):
    data, mapping = template(kind)
    preview = inspect(data, kind + '.xlsx', Mapping.model_validate(mapping))
    assert preview['total'] == 5
    assert [r['row'] for r in preview['rows']] == [header + 1, header + 2, header + 4, header + 5, 2]
    assert preview['rows'][0]['question'] == preview['rows'][1]['question']
    assert preview['rows'][0]['question_cell'] == f'B{header + 1}'
    assert preview['rows'][0]['answer_cell'] == f'C{header + 1}'
    assert preview['rows'][0]['source_identifier'] == 'Q1'
    assert preview['rows'][2]['writable'] is False
    questionnaire, _, _ = imported(service, kind)
    restarted = Questionnaires(Store(service.store.path))
    assert restarted.get('a', questionnaire['id'])['total'] == 5
    assert len(restarted.items('a', questionnaire['id'])['items']) == 5
    with service.store.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM workbook_assets').fetchone()[0] == 1


def properties(wb, targets):
    result = {'order': wb.sheetnames, 'sheets': {}}
    for ws in wb:
        result['sheets'][ws.title] = {
            'state': ws.sheet_state, 'merged': sorted(map(str, ws.merged_cells.ranges)),
            'widths': {k: v.width for k, v in ws.column_dimensions.items()},
            'rows': {k: (v.height, v.hidden) for k, v in ws.row_dimensions.items()},
            'validation': str(ws.data_validations), 'freeze': ws.freeze_panes,
            'cells': {cell.coordinate: (cell.value, cell.data_type, str(cell.font), str(cell.fill), str(cell.border), str(cell.alignment), cell.number_format)
                      for row in ws for cell in row if (ws.title, cell.coordinate) not in targets}}
    return result


@pytest.mark.parametrize('kind', ['security', 'rfp'])
@pytest.mark.parametrize('text', ['=HYPERLINK("https://invalid")', '+SUM(1,2)', '-1+2', '@SUM(1,2)'])
def test_roundtrip_approved_only_literal_source_immutable(service, kind, text):
    questionnaire, original, _ = imported(service, kind)
    service.draft('a', questionnaire['id'])
    items = service.items('a', questionnaire['id'])['items']
    first = service.edit('a', items[0]['id'], text, 1)
    service.approve('a', first['id'], 2, 'local reviewer')
    service.approve('a', items[1]['id'], 1, 'local reviewer')
    service.edit('a', items[1]['id'], 'Changed after approval', 1)
    service.approve('a', items[2]['id'], 1, 'local reviewer')  # existing cell is preserved
    output, manifest, unresolved = unpack(service.export('a', questionnaire['id']))
    targets = {(first['provenance']['sheet'], first['provenance']['answer_cell'])}
    before = load_workbook(io.BytesIO(original))
    after = load_workbook(io.BytesIO(output))
    assert properties(before, targets) == properties(after, targets)
    cell = after[first['provenance']['sheet']][first['provenance']['answer_cell']]
    assert cell.value == text and cell.data_type == 's'
    assert len(manifest['selected_revisions']) == 1
    assert len(manifest['unresolved_item_ids']) == 4
    assert 'existing_answer_preserved' in unresolved
    assert manifest['source_sha256'] == hashlib.sha256(original).hexdigest()
    assert manifest['output_sha256'] == hashlib.sha256(output).hexdigest()
    before.close()
    after.close()
    with service.store.connect() as db:
        assert bytes(db.execute('SELECT original FROM workbook_assets').fetchone()[0]) == original
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("UPDATE workbook_assets SET filename='changed.xlsx'")


def test_overwrite_hidden_intent_and_duplicate_mapping(service):
    data, mapping = template()
    mapping['sheets'][0]['overwrite_existing'] = True
    mapping['sheets'][0]['include_hidden'] = True
    mapping['sheets'].append({'sheet': 'Internal', 'header_row': 1, 'question_column': 'A', 'answer_column': 'B', 'include_hidden': True})
    questionnaire, _, _ = imported(service, raw_mapping=mapping)
    assert questionnaire['total'] == 7
    service.draft('a', questionnaire['id'])
    item = service.items('a', questionnaire['id'])['items'][2]
    service.approve('a', item['id'], 1, 'reviewer')
    output, manifest, _ = unpack(service.export('a', questionnaire['id']))
    wb = load_workbook(io.BytesIO(output))
    assert 'TLS' in wb['Security']['C5'].value
    assert manifest['selected_revisions'][0]['overwrite_existing'] is True
    wb.close()
    mapping['sheets'].append(mapping['sheets'][0])
    error('duplicate_destination', lambda: inspect(data, 'a.xlsx', Mapping.model_validate(mapping)))


@pytest.mark.parametrize('change,code', [
    (lambda wb: setattr(wb['Security']['C2'], 'value', '=1+2'), 'formula_destination'),
    (lambda wb: setattr(wb['Security']['B2'], 'value', '=1+2'), 'formula_question'),
    (lambda wb: wb['Security'].merge_cells('B2:C2'), 'merged_destination')])
def test_formula_and_merged_rejections(change, code):
    data, mapping = template()
    error(code, lambda: inspect(mutate(data, change), 'a.xlsx', Mapping.model_validate(mapping)))


def test_invalid_mapping_and_hidden_sheet():
    data, mapping = template()
    mapping['sheets'][0]['answer_column'] = 'B'
    error('invalid_mapping', lambda: inspect(data, 'a.xlsx', Mapping.model_validate(mapping)))
    mapping['sheets'] = [{'sheet': 'Internal', 'header_row': 1, 'question_column': 'A', 'answer_column': 'B'}]
    error('hidden_sheet', lambda: inspect(data, 'a.xlsx', Mapping.model_validate(mapping)))
    mapping['sheets'][0]['sheet'] = 'missing'
    error('invalid_mapping', lambda: inspect(data, 'a.xlsx', Mapping.model_validate(mapping)))
    mapping['sheets'][0]['sheet'] = 'Operations'
    mapping['sheets'].append({'sheet': 'Operations', 'header_row': 1, 'question_column': 'B', 'answer_column': 'C'})
    changed = mutate(data, lambda wb: setattr(wb['Operations']['B2'], 'value', 'Overlapping question'))
    error('overlapping_mapping', lambda: inspect(changed, 'a.xlsx', Mapping.model_validate(mapping)))


def test_idempotency_and_safe_filename(service):
    questionnaire, data, mapping = imported(service)
    assert service.create('a', data, 'security.xlsx', Mapping.model_validate(mapping), 'key')['id'] == questionnaire['id']
    mapping['sheets'][0]['overwrite_existing'] = True
    error('idempotency_conflict', lambda: service.create('a', data, 'security.xlsx', Mapping.model_validate(mapping), 'key'))
    other = service.create('b', data, '../../escape.xlsx', Mapping.model_validate(mapping), 'key')
    with service.store.connect() as db:
        asset = db.execute('SELECT * FROM workbook_assets WHERE id=?', (other['asset_id'],)).fetchone()
        assert asset['filename'] == 'escape.xlsx'
        assert '/' not in asset['storage_key'] and '\\' not in asset['storage_key']


def test_tenant_isolation_and_relations(service):
    questionnaire, _, _ = imported(service)
    service.draft('a', questionnaire['id'])
    item = service.items('a', questionnaire['id'])['items'][0]
    for operation in [lambda: service.get('b', questionnaire['id']), lambda: service.items('b', questionnaire['id']),
                      lambda: service.draft('b', questionnaire['id']), lambda: service.approve('b', item['id'], 1, 'reviewer'),
                      lambda: service.edit('b', item['id'], 'edit', 1), lambda: service.export('b', questionnaire['id'])]:
        error('not_found', operation)
    with service.store.connect() as db:
        with pytest.raises(sqlite3.IntegrityError):
            db.execute('UPDATE questionnaire_items SET tenant=? WHERE id=?', ('b', item['id']))


def test_concurrency_stale_approval_evidence_and_abstention(service):
    questionnaire, _, _ = imported(service)
    assert service.draft('a', questionnaire['id'], 2)['remaining'] == 3
    service.draft('a', questionnaire['id'])
    assert service.draft('a', questionnaire['id'])['drafted'] == 0
    items = service.items('a', questionnaire['id'])['items']
    service.approve('a', items[0]['id'], 1, 'reviewer')
    def edit(text):
        try:
            return service.edit('a', items[0]['id'], text, 1)['current_revision']
        except Problem as exc:
            return exc.code
    with ThreadPoolExecutor(2) as executor:
        results = list(executor.map(edit, ['edit one', 'edit two']))
    assert sorted(map(str, results)) == ['2', 'revision_conflict']
    error('revision_conflict', lambda: service.approve('a', items[0]['id'], 1, 'reviewer'))
    error('stale_evidence', lambda: service.approve('a', items[3]['id'], 1, 'reviewer'))
    service.approve('a', items[4]['id'], 1, 'reviewer')
    with service.store.connect() as db:
        db.execute("UPDATE documents SET version='2' WHERE tenant='a'")
    _, manifest, report = unpack(service.export('a', questionnaire['id']))
    assert manifest['selected_revisions'] == []
    assert 'stale_evidence' in report


def zip_changes(data, extra=None, rewrite=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            target.writestr(entry.filename, rewrite(entry.filename, content) if rewrite else content)
        if extra:
            target.writestr(*extra)
    return output.getvalue()


@pytest.mark.parametrize('extension', ['a.xls', 'a.xlsm', 'a.xltx', 'a.zip'])
def test_unsupported_extensions(extension):
    data, _ = template()
    error('unsupported_format', lambda: workbook(data, extension))


@pytest.mark.parametrize('code,kwargs', [
    ('upload_too_large', {'upload': 1}), ('archive_too_large', {'expanded': 1}),
    ('worksheet_too_large', {'sheets': 1}), ('worksheet_too_large', {'cells': 1}),
    ('worksheet_too_large', {'rows': 2}), ('archive_too_large', {'entries': 1})])
def test_limits_before_expensive_loader(code, kwargs):
    data, _ = template()
    with patch('app.workbooks.load_workbook') as loader:
        error(code, lambda: workbook(data, 'a.xlsx', replace(Limits(), **kwargs)))
        loader.assert_not_called()


def test_archive_attacks_and_unsupported_features():
    data, _ = template()
    for content, code in [(b'not zip', 'malformed_workbook'),
                          (zip_changes(data, ('../outside.xml', '<x/>')), 'unsafe_archive'),
                          (zip_changes(data, ('xl/drawings/drawing1.xml', '<x/>')), 'unsupported_feature'),
                          (zip_changes(data, ('xl/externalLinks/externalLink1.xml', '<x/>')), 'unsupported_feature')]:
        with patch('app.workbooks.load_workbook') as loader:
            error(code, lambda: workbook(content, 'a.xlsx'))
            loader.assert_not_called()
    def external(name, content):
        if name == 'xl/_rels/workbook.xml.rels':
            return content.replace(b'</Relationships>', b'<Relationship Id="rId99" Type="externalLink" Target="https://invalid" TargetMode="External"/></Relationships>')
        return content
    error('external_links', lambda: workbook(zip_changes(data, rewrite=external), 'a.xlsx'))
    def xml_entity(name, content):
        return b'<!DOCTYPE x [<!ENTITY y "x">]><x>&y;</x>' if name == 'xl/workbook.xml' else content
    error('malformed_workbook', lambda: workbook(zip_changes(data, rewrite=xml_entity), 'a.xlsx'))


def test_question_count_limit():
    data, mapping = template()
    error('too_many_questions', lambda: inspect(data, 'a.xlsx', Mapping.model_validate(mapping), replace(Limits(), questions=1)))


def test_encryption_large_merge_extensions_and_xml_errors():
    data, _ = template()
    encrypted = bytearray(data)
    central = encrypted.index(b'PK\x01\x02')
    struct.pack_into('<H', encrypted, central + 8, struct.unpack_from('<H', encrypted, central + 8)[0] | 1)
    with patch('app.workbooks.load_workbook') as loader:
        error('unsupported_format', lambda: workbook(bytes(encrypted), 'encrypted.xlsx'))
        loader.assert_not_called()
    def huge_merge(name, content):
        if name == 'xl/worksheets/sheet1.xml':
            return content.replace(b'C3:D3', b'C3:XFD10000')
        return content
    with patch('app.workbooks.load_workbook') as loader:
        error('worksheet_too_large', lambda: workbook(zip_changes(data, rewrite=huge_merge), 'big.xlsx'))
        loader.assert_not_called()
    def extensions(name, content):
        if name == 'xl/worksheets/sheet1.xml':
            return content.replace(b'</worksheet>', b'<extLst/></worksheet>')
        return content
    error('unsupported_feature', lambda: workbook(zip_changes(data, rewrite=extensions), 'ext.xlsx'))
    def bad_xml(name, content):
        return b'<unclosed>' if name == 'xl/workbook.xml' else content
    error('malformed_workbook', lambda: workbook(zip_changes(data, rewrite=bad_xml), 'bad.xlsx'))


def test_preview_has_no_database_side_effects_and_shipped_samples(service):
    before = {}
    with service.store.connect() as db:
        for table in ['workbook_assets', 'questionnaires', 'questionnaire_items', 'answers']:
            before[table] = db.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0]
    for kind in ['security', 'rfp']:
        samples = Path(__file__).resolve().parents[1] / 'samples'
        mapping = Mapping.model_validate_json((samples / f'{kind}.mapping.json').read_text())
        assert inspect((samples / f'{kind}.xlsx').read_bytes(), kind + '.xlsx', mapping)['total'] == 5
    with service.store.connect() as db:
        for table, count in before.items():
            assert db.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0] == count


@pytest.mark.parametrize('value', ['bad\x00answer', '\ud800', '\uffff', '   '])
def test_invalid_answer_text(service, value):
    questionnaire, _, _ = imported(service)
    service.draft('a', questionnaire['id'])
    item = service.items('a', questionnaire['id'])['items'][0]
    error('invalid_answer', lambda: service.edit('a', item['id'], value, 1))


@pytest.mark.parametrize('change', ["version='changed'", "valid_until='2000-01-01'", "content='changed content'"])
def test_each_evidence_change_excludes_approved_revision(service, change):
    questionnaire, _, _ = imported(service)
    service.draft('a', questionnaire['id'])
    item = service.items('a', questionnaire['id'])['items'][0]
    service.approve('a', item['id'], 1, 'reviewer')
    with service.store.connect() as db:
        db.execute(f"UPDATE documents SET {change} WHERE tenant='a'")
    _, manifest, report = unpack(service.export('a', questionnaire['id']))
    assert manifest['selected_revisions'] == []
    assert 'stale_evidence' in report


def test_migration_preserves_legacy_data(tmp_path):
    path = tmp_path / 'old.db'
    # Construct the exact pre-M1 schema without invoking the new Store initializer.
    db = sqlite3.connect(path)
    db.executescript('CREATE TABLE documents (id TEXT PRIMARY KEY,tenant TEXT NOT NULL,title TEXT NOT NULL,content TEXT NOT NULL,version TEXT NOT NULL,valid_until TEXT NOT NULL); CREATE TABLE answers (id TEXT PRIMARY KEY,tenant TEXT NOT NULL,question TEXT NOT NULL,draft TEXT NOT NULL,evidence TEXT NOT NULL,status TEXT NOT NULL,final TEXT,reviewer TEXT,trace TEXT NOT NULL);')
    db.execute('INSERT INTO documents VALUES (?,?,?,?,?,?)', ('doc', 'a', 'Policy', 'Legacy text', '1', '2099-01-01'))
    db.execute('INSERT INTO answers VALUES (?,?,?,?,?,?,?,?,?)', ('answer', 'a', 'Question', 'Draft', '[]', 'approved', 'Final', 'Reviewer', '{}'))
    db.commit()
    db.close()
    for _ in range(2):
        migrated = Store(path)
        assert migrated.get('a', 'answer')['final'] == 'Final'
        with migrated.connect() as db:
            assert db.execute('SELECT content FROM documents').fetchone()[0] == 'Legacy text'
            assert db.execute('SELECT COUNT(*) FROM schema_migrations').fetchone()[0] == 3


def test_api_upload_to_export_pagination_and_errors(tmp_path, monkeypatch):
    monkeypatch.setenv('EVIDENCEDESK_DB', str(tmp_path / 'api.db'))
    monkeypatch.setenv('EVIDENCEDESK_API_KEY', 'test-secret')
    monkeypatch.setenv('EVIDENCEDESK_TENANT', 'a')
    api_store.cache_clear()
    client = TestClient(app)
    headers = {'X-API-Key': 'test-secret', 'Idempotency-Key': 'api-import'}
    data, mapping = template()
    def upload(path, mapping_value=mapping, content=data):
        return client.post(path, headers=headers, files={'file': ('../../template.xlsx', content)}, data={'mapping': json.dumps(mapping_value)} if mapping_value is not None else {})
    try:
        assert upload('/questionnaires/preview', None).json()['mapping'] is None
        assert upload('/questionnaires/preview').json()['total'] == 5
        assert upload('/questionnaires', None).json()['error']['code'] == 'mapping_required'
        assert upload('/questionnaires/preview', {}, b'bad').status_code == 422
        assert client.post('/questionnaires', files={'file': ('x.xlsx', data)}).status_code == 401
        response = upload('/questionnaires')
        assert response.status_code == 201, response.text
        identifier = response.json()['id']
        assert upload('/questionnaires').json()['id'] == identifier
        client.post('/documents', headers=headers, json={'title': 'Policy', 'content': 'Customer data is encrypted at rest using AES-256.', 'version': '1', 'valid_until': '2099-01-01'})
        assert client.post(f'/questionnaires/{identifier}/draft', headers=headers, json={'limit': 2}).json()['remaining'] == 3
        page = client.get(f'/questionnaires/{identifier}/items?limit=2', headers=headers).json()
        assert page['total'] == 5 and page['next_offset'] == 2
        assert client.get(f'/questionnaires/{identifier}/items?offset=4&limit=2', headers=headers).json()['next_offset'] is None
        item = page['items'][0]
        path = f"/questionnaire-items/{item['id']}"
        assert client.patch(path + '/answer', headers=headers, json={'text': '=literal', 'expected_revision': 1, 'owner': 'SME'}).status_code == 200
        assert client.post(path + '/approve', headers=headers, json={'expected_revision': 1, 'reviewer': 'local'}).json()['error']['code'] == 'revision_conflict'
        assert client.post(path + '/approve', headers=headers, json={'expected_revision': 2, 'reviewer': 'local'}).status_code == 200
        response = client.post(f'/questionnaires/{identifier}/export', headers=headers)
        assert response.status_code == 200
        _, manifest, _ = unpack(response.content)
        assert len(manifest['selected_revisions']) == 1
        assert client.get(f'/questionnaires/{identifier}/items?limit=0', headers=headers).json()['error']['code'] == 'invalid_pagination'
        assert client.get(f'/questionnaires/{identifier}/items?limit=x', headers=headers).json()['error']['code'] == 'invalid_request'
        monkeypatch.setenv('EVIDENCEDESK_TENANT', 'b')
        for method, path_value in [('get', f'/questionnaires/{identifier}'), ('get', f'/questionnaires/{identifier}/items'), ('post', f'/questionnaires/{identifier}/export')]:
            assert getattr(client, method)(path_value, headers=headers).status_code == 404
        assert client.post(path + '/approve', headers=headers, json={'expected_revision': 2, 'reviewer': 'other'}).status_code == 404
        monkeypatch.setenv('EVIDENCEDESK_MAX_UPLOAD', '10')
        assert upload('/questionnaires/preview').json()['error']['code'] == 'upload_too_large'
    finally:
        api_store.cache_clear()

def test_dashboard_pagination_isolation_and_restart(service):
    first, _, _ = imported(service, key='one')
    second, _, _ = imported(service, key='two')
    data, mapping = template('security')
    service.create('b', data, 'private.xlsx', Mapping.model_validate(mapping), 'other')
    page = service.list('a', limit=1)
    assert page['total'] == 2 and page['next_offset'] == 1
    assert set(page['items'][0]) == {'id', 'filename', 'status', 'created_at', 'total'}
    assert page['items'][0]['total'] == 5
    next_page = service.list('a', offset=1, limit=1)
    assert next_page['next_offset'] is None
    assert {page['items'][0]['id'], next_page['items'][0]['id']} == {first['id'], second['id']}
    assert service.list('unknown')['items'] == []
    assert service.list('a', offset=100)['items'] == []
    reopened = Questionnaires(Store(service.store.path))
    assert reopened.list('a')['total'] == 2
    for offset, limit in [(-1, 1), (0, 0), (0, 101)]:
        error('invalid_pagination', lambda: service.list('a', offset, limit))


def test_dashboard_http_pagination(service, monkeypatch):
    imported(service)
    monkeypatch.setenv("EVIDENCEDESK_DB", service.store.path)
    api_store.cache_clear()
    app.dependency_overrides[tenant] = lambda: 'a'
    try:
        with TestClient(app) as client:
            assert client.get('/questionnaires?limit=1').json()['total'] == 1
            assert client.get('/questionnaires?offset=-1').status_code == 422
            assert client.get('/questionnaires?limit=101').status_code == 422
            assert client.get('/questionnaires?offset=bad').status_code == 422
            app.dependency_overrides[tenant] = lambda: 'b'
            assert client.get('/questionnaires').json()['total'] == 0
    finally:
        app.dependency_overrides.clear()
        api_store.cache_clear()


def test_workbook_limits_authenticated_and_configured(monkeypatch):
    monkeypatch.setenv('EVIDENCEDESK_API_KEY', 'limit-test')
    monkeypatch.setenv('EVIDENCEDESK_MAX_UPLOAD', '1234')
    with TestClient(app) as client:
        assert client.get('/questionnaires/limits').status_code == 401
        response = client.get('/questionnaires/limits', headers={'X-API-Key': 'limit-test'})
        assert response.status_code == 200
        assert response.json() == {'max_upload_bytes': 1234}

def test_item_read_api_tenant_isolation(service, monkeypatch):
    questionnaire, _, _ = imported(service)
    identifier = service.items('a', questionnaire['id'])['items'][0]['id']
    monkeypatch.setenv('EVIDENCEDESK_DB', service.store.path)
    monkeypatch.setenv('EVIDENCEDESK_API_KEY', 'read-test')
    monkeypatch.setenv('EVIDENCEDESK_TENANT', 'a')
    api_store.cache_clear()
    try:
        with TestClient(app) as client:
            assert client.get('/questionnaire-items/' + identifier).status_code == 401
            response = client.get('/questionnaire-items/' + identifier, headers={'X-API-Key': 'read-test'})
            assert response.status_code == 200
            assert response.json()['current_revision'] == 0
            monkeypatch.setenv('EVIDENCEDESK_TENANT', 'b')
            assert client.get('/questionnaire-items/' + identifier, headers={'X-API-Key': 'read-test'}).status_code == 404
    finally:
        api_store.cache_clear()

def test_export_receipt_matches_bundle_and_survives_later_edits(service):
    questionnaire, _, _ = imported(service)
    identifier = questionnaire['id']
    service.draft('a', identifier)
    items = service.items('a', identifier)['items']
    service.approve('a', items[0]['id'], 1, 'local')
    service.edit('a', items[3]['id'], 'Needs a specialist.', 1, owner='Evidence owner')
    bundle, manifest = service.export_result('a', identifier)
    _, bundled_manifest, report = unpack(bundle)
    assert bundled_manifest == manifest
    summary = service.export_summary('a', identifier, manifest['export_id'], limit=1)
    assert summary['selected_count'] == 1 and summary['unresolved_count'] == 4 and summary['total'] == 5
    assert summary['details_available'] and summary['next_offset'] == 1
    assert len(summary['items']) == 1
    all_rows = service.export_summary('a', identifier, manifest['export_id'])['items']
    assert all_rows == manifest['unresolved_items']
    assert 'Evidence owner' in report
    service.edit('a', items[0]['id'], 'Edited after export.', 1)
    service.edit('a', items[3]['id'], 'Later owner change.', 2, owner='New owner')
    assert service.export_summary('a', identifier, manifest['export_id'])['items'] == all_rows
    assert service.export_summary('a', identifier, manifest['export_id'])['selected_count'] == 1
    error('not_found', lambda: service.export_summary('b', identifier, manifest['export_id']))
    other, _, _ = imported(service, key='other-export')
    error('not_found', lambda: service.export_summary('a', other['id'], manifest['export_id']))
    for offset, limit in [(-1, 20), (0, 0), (0, 101)]:
        error('invalid_pagination', lambda: service.export_summary('a', identifier, manifest['export_id'], offset, limit))
    assert service.export_summary('a', identifier, manifest['export_id'], offset=99)['items'] == []


def test_historical_export_receipt_does_not_invent_old_details(service):
    questionnaire, _, _ = imported(service)
    _, manifest = service.export_result('a', questionnaire['id'])
    manifest.pop('unresolved_items')
    with service.store.connect() as db:
        db.execute('UPDATE questionnaire_exports SET manifest=? WHERE id=?', (json.dumps(manifest), manifest['export_id']))
    summary = service.export_summary('a', questionnaire['id'], manifest['export_id'])
    assert summary['unresolved_count'] == 5 and summary['selected_count'] == 0
    assert not summary['details_available']
    assert all(row['reason'] == 'details_unavailable' and row['owner'] is None for row in summary['items'])


def test_export_receipt_http_header_auth_scope_and_pagination(service, monkeypatch):
    questionnaire, _, _ = imported(service)
    identifier = questionnaire['id']
    monkeypatch.setenv('EVIDENCEDESK_DB', service.store.path)
    monkeypatch.setenv('EVIDENCEDESK_API_KEY', 'export-test')
    monkeypatch.setenv('EVIDENCEDESK_TENANT', 'a')
    api_store.cache_clear()
    headers = {'X-API-Key': 'export-test'}
    try:
        with TestClient(app) as client:
            response = client.post(f'/questionnaires/{identifier}/export', headers=headers)
            assert response.status_code == 200 and response.headers['content-type'] == 'application/zip'
            _, manifest, _ = unpack(response.content)
            assert response.headers['x-export-id'] == manifest['export_id']
            assert response.headers['cache-control'] == 'no-store'
            path = f"/questionnaires/{identifier}/exports/{manifest['export_id']}"
            assert client.get(path).status_code == 401
            summary = client.get(path + '?limit=2', headers=headers).json()
            assert len(summary['items']) == 2 and summary['next_offset'] == 2
            assert set(summary) == {'export_id', 'questionnaire_id', 'created_at', 'total', 'selected_count', 'unresolved_count', 'items', 'offset', 'limit', 'details_available', 'next_offset'}
            assert client.get(path + '?offset=-1', headers=headers).status_code == 422
            assert client.get(path + '?limit=101', headers=headers).status_code == 422
            assert client.get(path + '?limit=bad', headers=headers).status_code == 422
            monkeypatch.setenv('EVIDENCEDESK_TENANT', 'b')
            assert client.get(path, headers=headers).status_code == 404
    finally:
        api_store.cache_clear()
